"""
Generates migration_toolset_guide.xlsx — folder structure + step-by-step process guide.
Run once: python create_folder_guide.py
"""

import openpyxl
from openpyxl.styles import (Font, PatternFill, Alignment, Border, Side,
                              GradientFill)
from openpyxl.utils import get_column_letter

# ─────────────────────────────────────────────────────────────────────────────
# Content
# ─────────────────────────────────────────────────────────────────────────────

# (step, category, path, description)
# step = '' for folder/output rows (not a numbered tool step)
ROWS = [
    # ── PRE-REQUISITES ────────────────────────────────────────────────────────
    ('',   'ROOT',       'migration_toolset\\',
     'Root folder for the entire Toad → Azure migration toolset'),

    ('',   'INPUT',      'migration_toolset\\input\\',
     'Place raw Toad XML automation files here before running Tool 1'),

    ('',   'GLOBAL',     'migration_toolset\\global\\',
     'Shared resources used across all 200+ reports'),

    ('',   'GLOBAL',     'migration_toolset\\global\\mapping_registry.csv',
     'Global registry: maps every real value (email/server/password) to its mock replacement. GITIGNORED.'),

    ('',   'GLOBAL',     'migration_toolset\\global\\poc_team_config.json',
     'Dev team email addresses for --poc testing mode. Fill before running Tool 3 --poc.'),

    ('',   'GLOBAL',     'migration_toolset\\global\\actual_mapping_files\\{report}_actual_mapping.csv',
     'Per-report actual email list (real client emails). GITIGNORED — never commit.'),

    ('',   'GLOBAL',     'migration_toolset\\global\\poc_mapping_files\\{report}_poc_mapping.csv',
     'Per-report POC email list (dev team emails mapped to each report). Safe to commit.'),

    ('',   'GLOBAL',     'migration_toolset\\global\\final_actual_mapping.csv',
     'Consolidated: DISTINCT (report, real_email) across all processed reports. GITIGNORED.'),

    ('',   'GLOBAL',     'migration_toolset\\global\\final_poc_mapping.csv',
     'Consolidated: every report x dev team emails. Used for bulk Tool 3 --poc runs.'),

    ('',   'GLOBAL',     'migration_toolset\\global\\xlsm_templates\\{report_name}\\',
     'Place the report Excel template (.xlsm) here before running Tool 5 blob upload.'),

    ('',   'GLOBAL',     'migration_toolset\\global\\config\\config_schema_ddl.sql',
     'Tool 3 output: CREATE TABLE DDL for config schema (config.report, config.report_email, config.report_sql). Run once per environment.'),

    # ── STEP 1 ────────────────────────────────────────────────────────────────
    ('STEP 1', 'TOOL',   'migration_toolset\\tools\\tool1_sanitize.py',
     'Sanitize Toad XML: replace all sensitive values (emails, passwords, servers, UNC paths) with safe mock values'),

    ('',   'OUTPUT',     'migration_toolset\\reports\\{report}\\sanitized\\{report}_sanitized.txt',
     'Sanitized XML file — safe to share, use in all downstream tools'),

    ('',   'OUTPUT',     'migration_toolset\\reports\\{report}\\sanitized\\{report}_mapping.csv',
     'Per-report mapping: real value → mock value for this report only'),

    # ── STEP 2 (POC only) ─────────────────────────────────────────────────────
    ('STEP 2\n(POC only)', 'TOOL', 'migration_toolset\\tools\\tool1_01_mock_data.py',
     'Mock Data Generator: reads 10 seed rows per table from input\\seed\\ → generates 1000+ synthetic rows'),

    ('',   'INPUT',      'migration_toolset\\input\\seed\\{report}\\{table}.csv',
     'User-supplied: 10 real rows per table (anonymised). Used as seed for mock data generation.'),

    ('',   'OUTPUT',     'migration_toolset\\reports\\{report}\\mock_data\\{table}.csv',
     'Generated mock data: 10 seed rows + 1000 synthetic rows per table'),

    # ── STEP 3 (POC only) ─────────────────────────────────────────────────────
    ('STEP 3\n(POC only)', 'TOOL', 'migration_toolset\\tools\\tool1_02_ddl_generator.py',
     'DDL Generator: infers SQL column types from mock CSVs + schema from sanitized XML → generates CREATE TABLE scripts'),

    ('',   'OUTPUT',     'migration_toolset\\reports\\{report}\\ddl\\{report}_ddl.sql',
     'T-SQL CREATE TABLE DDL for all tables in the report. Idempotent (IF OBJECT_ID guard).'),

    # ── STEP 4 ────────────────────────────────────────────────────────────────
    ('STEP 4', 'TOOL',   'migration_toolset\\tools\\tool2_adf_generator.py',
     'ADF ARM Template Generator: reads sanitized XML → generates Linked Services, Datasets, Pipeline JSON + combined ARM template'),

    ('',   'OUTPUT',     'migration_toolset\\reports\\{report}\\adf\\linkedServices\\',
     'ADF Linked Service JSONs: LS_AzureKeyVault, LS_AzureSQL_Source, LS_AzureBlobStorage'),

    ('',   'OUTPUT',     'migration_toolset\\reports\\{report}\\adf\\datasets\\',
     'ADF Dataset JSONs: one per SQL table + ODS log + report source + blob output/archive + config'),

    ('',   'OUTPUT',     'migration_toolset\\reports\\{report}\\adf\\pipelines\\PL_{report}.json',
     'ADF Pipeline JSON: full activity chain (ODS check → set vars → copy data → email notifications)'),

    ('',   'OUTPUT',     'migration_toolset\\reports\\{report}\\adf\\arm_template.json',
     'Combined deployable ARM template: all 16 resources (factory + LS + datasets + pipeline)'),

    ('',   'OUTPUT',     'migration_toolset\\reports\\{report}\\adf\\arm_template_parameters.json',
     'Parameter template: fill server names, Key Vault name, storage account before deploying. GITIGNORED.'),

    # ── STEP 5 ────────────────────────────────────────────────────────────────
    ('STEP 5', 'TOOL',   'migration_toolset\\tools\\tool3_config_setup.py',
     'Config Setup: reads sanitized XML → generates config schema DDL + per-report INSERT SQL (emails, SQL queries, template info)'),

    ('',   'OUTPUT',     'migration_toolset\\reports\\{report}\\config\\{report}_config_data.sql',
     'Idempotent INSERT SQL for config.report (+ template), config.report_sql (ODS_CHECK + REPORT_MAIN), config.report_email'),

    # ── STEP 6 ────────────────────────────────────────────────────────────────
    ('STEP 6', 'TOOL',   'migration_toolset\\tools\\tool4_adf_deploy.py',
     'ADF Deploy Script Generator: splits ARM template into bootstrap (once) + per-report; generates PS1 + bash deploy scripts'),

    ('',   'OUTPUT',     'migration_toolset\\reports\\{report}\\adf\\deploy\\deploy_bootstrap.ps1 / .sh',
     'Run once per environment: deploys ADF factory + 3 shared Linked Services'),

    ('',   'OUTPUT',     'migration_toolset\\reports\\{report}\\adf\\deploy\\deploy_{report}.ps1 / .sh',
     'Run per report: deploys 11 datasets + 1 pipeline (Incremental mode)'),

    ('',   'OUTPUT',     'migration_toolset\\reports\\{report}\\adf\\deploy\\deployment_checklist.txt',
     'Parameter status checklist: shows [OK]/[!!] per parameter + full resource list + run order'),

    # ── STEP 7 ────────────────────────────────────────────────────────────────
    ('STEP 7', 'TOOL',   'migration_toolset\\tools\\tool5_blob_upload.py',
     'Blob Upload Script Generator: extracts xlsm template filename from sanitized XML → generates PS1 + bash az storage upload scripts'),

    ('',   'OUTPUT',     'migration_toolset\\reports\\{report}\\blob\\upload_{report}.ps1 / .sh',
     'Upload script: az storage blob upload for xlsm template → templates/{report}/{file}.xlsm in Blob Storage'),

    ('',   'OUTPUT',     'migration_toolset\\reports\\{report}\\blob\\blob_checklist.txt',
     'Template presence check: [OK] if xlsm found in global\\xlsm_templates\\{report}\\, [!!] if missing'),

    # ── UTILITIES ─────────────────────────────────────────────────────────────
    ('',   'UTILITY',    'migration_toolset\\tools\\utils\\xml_parser.py',
     'ToadXmlParser: extracts all sensitive values from Toad XML (emails, passwords, servers, UNC paths, SQL tables)'),

    ('',   'UTILITY',    'migration_toolset\\tools\\utils\\mapping_registry.py',
     'MappingRegistry: loads/saves global CSV, generates consistent mock replacements across all reports'),

    # ── STREAMLIT UI ──────────────────────────────────────────────────────────
    ('UI',  'UI',        'migration_toolset\\app.py',
     'Streamlit web UI: run all 5 tools from a browser without using the command line. '
     'Run: streamlit run app.py  (from migration_toolset\\ directory)'),

    ('',    'UI',        'http://localhost:8501',
     'Browser address once app.py is running. Streamlit auto-opens on first launch.'),
]

# ─────────────────────────────────────────────────────────────────────────────
# Colour palette
# ─────────────────────────────────────────────────────────────────────────────

COL_HEADER   = '1F4E79'   # dark blue
COL_STEP     = '2E75B6'   # medium blue
COL_TOOL     = 'D6E4F0'   # light blue
COL_OUTPUT   = 'EBF3E8'   # light green
COL_INPUT    = 'FFF2CC'   # light yellow
COL_GLOBAL   = 'F2F2F2'   # light grey
COL_ROOT     = 'D9D9D9'   # grey
COL_UTILITY  = 'EDE7F6'   # light purple
COL_UI       = 'FCE4EC'   # light pink

CATEGORY_FILL = {
    'TOOL':    COL_TOOL,
    'OUTPUT':  COL_OUTPUT,
    'INPUT':   COL_INPUT,
    'GLOBAL':  COL_GLOBAL,
    'ROOT':    COL_ROOT,
    'UTILITY': COL_UTILITY,
    'UI':      COL_UI,
}

def _fill(hex_color):
    return PatternFill('solid', fgColor=hex_color)

def _border():
    thin = Side(style='thin', color='BFBFBF')
    return Border(left=thin, right=thin, top=thin, bottom=thin)


# ─────────────────────────────────────────────────────────────────────────────
# Build workbook
# ─────────────────────────────────────────────────────────────────────────────

wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'Migration Toolset Guide'

# ── Column widths ─────────────────────────────────────────────────────────────
ws.column_dimensions['A'].width = 14   # Step
ws.column_dimensions['B'].width = 72   # Path
ws.column_dimensions['C'].width = 70   # Description

# ── Header row ────────────────────────────────────────────────────────────────
headers = ['Step', 'Folder / File Path', 'Description']
for col, h in enumerate(headers, start=1):
    cell = ws.cell(row=1, column=col, value=h)
    cell.font      = Font(bold=True, color='FFFFFF', size=12)
    cell.fill      = _fill(COL_HEADER)
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    cell.border    = _border()

ws.row_dimensions[1].height = 24

# ── Data rows ─────────────────────────────────────────────────────────────────
for i, (step, category, path, desc) in enumerate(ROWS, start=2):
    bg = CATEGORY_FILL.get(category, 'FFFFFF')
    is_step = bool(step)

    # Step column
    c_step = ws.cell(row=i, column=1, value=step)
    c_step.font      = Font(bold=is_step, color=('FFFFFF' if is_step else '595959'), size=10)
    c_step.fill      = _fill(COL_STEP if is_step else bg)
    c_step.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    c_step.border    = _border()

    # Path column
    c_path = ws.cell(row=i, column=2, value=path)
    c_path.font      = Font(bold=is_step, name='Courier New', size=9,
                            color=('1F4E79' if is_step else '2C2C2C'))
    c_path.fill      = _fill(bg)
    c_path.alignment = Alignment(vertical='center', wrap_text=True)
    c_path.border    = _border()

    # Description column
    c_desc = ws.cell(row=i, column=3, value=desc)
    c_desc.font      = Font(size=10, color='2C2C2C')
    c_desc.fill      = _fill(bg)
    c_desc.alignment = Alignment(vertical='center', wrap_text=True)
    c_desc.border    = _border()

    ws.row_dimensions[i].height = 36 if '\n' in step else 28

# ── Freeze top row ────────────────────────────────────────────────────────────
ws.freeze_panes = 'A2'

# ── Legend sheet ──────────────────────────────────────────────────────────────
wl = wb.create_sheet('Legend')
wl.column_dimensions['A'].width = 18
wl.column_dimensions['B'].width = 55

legend_header = wl.cell(row=1, column=1, value='Colour Legend')
legend_header.font  = Font(bold=True, color='FFFFFF', size=12)
legend_header.fill  = _fill(COL_HEADER)
wl.merge_cells('A1:B1')
wl.cell(row=1, column=1).alignment = Alignment(horizontal='center')

legend_rows = [
    (COL_STEP,    'STEP N',   'Numbered tool step — run in this order'),
    (COL_TOOL,    'TOOL',     'Python tool script to run'),
    (COL_INPUT,   'INPUT',    'Files you must provide before running the step'),
    (COL_OUTPUT,  'OUTPUT',   'Files generated automatically by the tool'),
    (COL_GLOBAL,  'GLOBAL',   'Shared files used across all reports'),
    (COL_ROOT,    'ROOT',     'Root folder'),
    (COL_UTILITY, 'UTILITY',  'Shared utility modules (used by tools internally)'),
    (COL_UI,      'UI',       'Streamlit web UI — browser-based interface for all tools'),
]

for r, (color, label, meaning) in enumerate(legend_rows, start=2):
    ca = wl.cell(row=r, column=1, value=label)
    ca.fill      = _fill(color)
    ca.font      = Font(bold=True, size=10)
    ca.alignment = Alignment(horizontal='center', vertical='center')
    ca.border    = _border()
    cb = wl.cell(row=r, column=2, value=meaning)
    cb.fill      = _fill(color)
    cb.font      = Font(size=10)
    cb.alignment = Alignment(vertical='center')
    cb.border    = _border()
    wl.row_dimensions[r].height = 22

# ─────────────────────────────────────────────────────────────────────────────
# Streamlit UI sheet
# ─────────────────────────────────────────────────────────────────────────────

wu = wb.create_sheet('Streamlit UI')
wu.column_dimensions['A'].width = 22
wu.column_dimensions['B'].width = 30
wu.column_dimensions['C'].width = 65

# Header
for col, title in enumerate(['Section', 'Feature / Tab', 'Description'], start=1):
    cell = wu.cell(row=1, column=col, value=title)
    cell.font      = Font(bold=True, color='FFFFFF', size=12)
    cell.fill      = _fill(COL_HEADER)
    cell.alignment = Alignment(horizontal='center', vertical='center')
    cell.border    = _border()
wu.row_dimensions[1].height = 24

UI_ROWS = [
    # ── How to run ───────────────────────────────────────────────────────────
    ('HOW TO RUN',   'Start the app',
     'cd migration_toolset\\   then:   streamlit run app.py'),
    ('',             'Browser address',
     'http://localhost:8501  (auto-opens on first launch)'),
    ('',             'Stop the app',
     'Press Ctrl+C in the terminal where streamlit is running'),

    # ── Sidebar ──────────────────────────────────────────────────────────────
    ('SIDEBAR',      'Run Mode toggle',
     'Switch between Single Report (one report at a time) and Batch Run (multiple files)'),
    ('',             'Single — Choose existing',
     'Dropdown of all folders already in reports\\ — select to set active report'),
    ('',             'Single — Enter name',
     'Type a new report name manually (used before Tool 1 creates the folder)'),
    ('',             'Single — Progress',
     'Live step indicators: Step 1-5 shown as Done (green) or Pending (red) based on output files'),
    ('',             'Batch — Select first N',
     'Number input: type any count 1 to total, click Select to queue first N files from input\\'),
    ('',             'Batch — Select All',
     'One click to queue every file in input\\'),
    ('',             'Batch — Multiselect',
     'Manually add or remove individual files from the queue'),

    # ── Single Report tabs ────────────────────────────────────────────────────
    ('SINGLE REPORT\nTab 1 — Sanitize',
     'Upload file',
     'Upload a new Toad XML file (.txt or .xml) — saved automatically to input\\'),
    ('',             'Select existing',
     'Pick a file already in input\\ from a dropdown'),
    ('',             'Run Sanitize',
     'Calls Tool 1: replaces all sensitive values, writes sanitized XML + mapping CSV'),
    ('',             'Output',
     'Shows Tool 1 console output; mapping table (real → mock); download buttons for sanitized XML + mapping CSV'),

    ('SINGLE REPORT\nTab 2 — ADF Templates',
     'Prerequisite check',
     'Warns if sanitized XML is missing — must run Tab 1 first'),
    ('',             'Run Generate',
     'Calls Tool 2: generates Linked Services, Datasets, Pipeline JSON, ARM template'),
    ('',             'Output',
     'Console output; expandable list of all generated JSON files; download arm_template.json + parameters.json'),

    ('SINGLE REPORT\nTab 3 — Config Setup',
     'POC mode toggle',
     'Uses dev team emails from poc_team_config.json instead of mock emailid@company.com addresses'),
    ('',             'Run Generate',
     'Calls Tool 3: generates config schema DDL + per-report INSERT SQL (report, emails, SQL queries, template)'),
    ('',             'Output',
     'Console output; SQL preview expanders; download config_data.sql + config_schema_ddl.sql'),

    ('SINGLE REPORT\nTab 4 — Deploy Scripts',
     'Resource Group input',
     'Optional: enter Azure resource group name — defaults to <your-resource-group> placeholder'),
    ('',             'Run Generate',
     'Calls Tool 4: splits ARM template into bootstrap + per-report; generates PS1 + bash scripts'),
    ('',             'Output',
     'Console output; deployment checklist expander (OK/!! per parameter); download 4 scripts'),

    ('SINGLE REPORT\nTab 5 — Blob Upload',
     'Upload .xlsm template',
     'Upload the Excel report template — saved to global\\xlsm_templates\\{report}\\'),
    ('',             'Run Generate',
     'Calls Tool 5: detects template filename from XML, generates PS1 + bash upload scripts'),
    ('',             'Output',
     'Console output; blob checklist expander; download upload PS1 + SH scripts'),

    ('SINGLE REPORT\nTab 6 — POC Tools',
     'Tool 1.01 — Mock Data',
     'Set row count (100–10000), run generator, preview any table as a dataframe (POC only)'),
    ('',             'Tool 1.02 — DDL',
     'Run DDL generator (requires mock data), preview SQL in expander, download DDL file (POC only)'),

    # ── Batch Run ─────────────────────────────────────────────────────────────
    ('BATCH RUN',    'Tool checkboxes',
     'Choose which tools to run per report: T1+T2+T3 ticked by default, T4+T5 unticked'),
    ('',             'Options',
     'POC mode (Tool 3), Verbose output, Resource Group name (Tool 4)'),
    ('',             'Run Batch button',
     'Processes each selected file in order — runs chosen tools sequentially per report'),
    ('',             'Progress bar',
     'Updates file-by-file showing current file name and count (e.g. Processing 3/10)'),
    ('',             'Results table',
     'One row per file: File, Report Name, T1-T5 status (green tick / red cross / dash if skipped)'),
    ('',             'Summary metrics',
     'Total files processed / All steps OK count / Files with errors count'),
    ('',             'Clear results',
     'Resets the results table so you can run again'),
]

COL_SECTION  = 'C5E1F5'   # light sky
COL_BLANK    = 'FFFFFF'
COL_HOW      = 'E8F5E9'   # light green
COL_SINGLE   = 'FFF8E1'   # light amber
COL_BATCH    = 'FCE4EC'   # light pink

def _ui_bg(section):
    if section == 'HOW TO RUN':   return COL_HOW
    if section.startswith('SINGLE'): return COL_SINGLE
    if section == 'BATCH RUN':    return COL_BATCH
    if section == 'SIDEBAR':      return COL_SECTION
    return COL_BLANK

current_section = ''
for i, (section, feature, desc) in enumerate(UI_ROWS, start=2):
    if section:
        current_section = section
    bg = _ui_bg(current_section)

    ca = wu.cell(row=i, column=1, value=section)
    ca.font      = Font(bold=bool(section), size=10,
                        color='1F4E79' if section else '595959')
    ca.fill      = _fill(bg)
    ca.alignment = Alignment(vertical='center', wrap_text=True)
    ca.border    = _border()

    cb = wu.cell(row=i, column=2, value=feature)
    cb.font      = Font(bold=True, size=10)
    cb.fill      = _fill(bg)
    cb.alignment = Alignment(vertical='center', wrap_text=True)
    cb.border    = _border()

    cc = wu.cell(row=i, column=3, value=desc)
    cc.font      = Font(size=10)
    cc.fill      = _fill(bg)
    cc.alignment = Alignment(vertical='center', wrap_text=True)
    cc.border    = _border()

    wu.row_dimensions[i].height = 32 if '\n' in section else 26

wu.freeze_panes = 'A2'

# ─────────────────────────────────────────────────────────────────────────────
# Save
# ─────────────────────────────────────────────────────────────────────────────

import os
out = os.path.join(os.path.dirname(__file__), 'migration_toolset_guide.xlsx')
wb.save(out)
print(f'Created: {out}')
