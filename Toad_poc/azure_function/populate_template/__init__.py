"""
populate_template — Azure Function
Called by ADF after Copy_ReportData_To_Blob.

Expected POST body:
{
    "report_date": "13.05.2026",
    "csv_blob":    "output/BC_BIMIO_267_PRECOPY_13.05.2026.csv",   (optional override)
    "out_blob":    "output/BC_BIMIO_267_PRECOPY_13.05.2026.xlsm"   (optional override)
}

Steps:
  1. Download CSV from Blob Storage
  2. Download .xlsm template from Blob Storage
  3. Clear Raw Data sheet (keep row 1 headers)
  4. Write CSV rows into Raw Data sheet from row 2
  5. Upload populated .xlsm back to Blob Storage
  6. Return 200 OK with output blob path
"""

import os
import io
import logging
import json

import azure.functions as func
from azure.storage.blob import BlobServiceClient
import openpyxl
import pandas as pd

CONTAINER      = "toad-poc-reports"
TEMPLATE_BLOB  = "templates/BC_BIMIO_267_TEMPLATE.xlsm"
RAW_DATA_SHEET = "Raw Data"


def main(req: func.HttpRequest) -> func.HttpResponse:
    logging.info("populate_template — triggered")

    # ── Parse request body ─────────────────────────────────────────────────────
    try:
        body = req.get_json()
    except ValueError:
        return func.HttpResponse("Invalid JSON body", status_code=400)

    report_date = body.get("report_date")
    if not report_date:
        return func.HttpResponse("Missing 'report_date' in request body", status_code=400)

    csv_blob = body.get("csv_blob",  f"output/BC_BIMIO_267_PRECOPY_{report_date}.csv")
    out_blob = body.get("out_blob",  f"output/BC_BIMIO_267_PRECOPY_{report_date}.xlsm")

    logging.info(f"  report_date : {report_date}")
    logging.info(f"  csv_blob    : {csv_blob}")
    logging.info(f"  out_blob    : {out_blob}")

    # ── Connect to Blob Storage ────────────────────────────────────────────────
    conn_str = os.environ["STORAGE_CONNECTION_STRING"]
    client   = BlobServiceClient.from_connection_string(conn_str)
    container = client.get_container_client(CONTAINER)

    # ── Download CSV ───────────────────────────────────────────────────────────
    logging.info(f"  Downloading CSV: {csv_blob}")
    csv_bytes = container.get_blob_client(csv_blob).download_blob().readall()
    df = pd.read_csv(io.BytesIO(csv_bytes))
    logging.info(f"  CSV rows: {len(df)}")

    # ── Download template ──────────────────────────────────────────────────────
    logging.info(f"  Downloading template: {TEMPLATE_BLOB}")
    tmpl_bytes = container.get_blob_client(TEMPLATE_BLOB).download_blob().readall()
    wb = openpyxl.load_workbook(io.BytesIO(tmpl_bytes), keep_vba=True)
    ws = wb[RAW_DATA_SHEET]

    # ── Clear existing data rows (keep header row 1) ───────────────────────────
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.value = None

    # ── Write data rows from row 2 ─────────────────────────────────────────────
    for row_idx, row_data in enumerate(df.itertuples(index=False), start=2):
        for col_idx, value in enumerate(row_data, start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)

    logging.info(f"  Populated {len(df)} rows into '{RAW_DATA_SHEET}' sheet")

    # ── Upload populated workbook ──────────────────────────────────────────────
    out_stream = io.BytesIO()
    wb.save(out_stream)
    out_stream.seek(0)

    container.get_blob_client(out_blob).upload_blob(out_stream, overwrite=True)
    logging.info(f"  Uploaded: {out_blob}")

    return func.HttpResponse(
        json.dumps({"status": "success", "output_blob": out_blob}),
        status_code=200,
        mimetype="application/json"
    )
