"""
Creates a single-slide "one pager" PowerPoint for the Utilitics project.
Run: python create_ppt.py
Output: Utilitics_OnePager.pptx
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
import pptx.oxml.ns as nsmap
from lxml import etree

# ── Slide size: Widescreen 16:9 ──────────────────────────────────────────────
prs = Presentation()
prs.slide_width  = Inches(13.33)
prs.slide_height = Inches(7.5)

slide_layout = prs.slide_layouts[6]   # blank
slide = prs.slides.add_slide(slide_layout)

# ── Helper colours ────────────────────────────────────────────────────────────
DARK_BG    = RGBColor(0x0F, 0x11, 0x17)
PANEL_BG   = RGBColor(0x0D, 0x19, 0x29)
HEADER_BG  = RGBColor(0x1F, 0x6F, 0xEB)
GREEN      = RGBColor(0x23, 0x86, 0x36)
ORANGE     = RGBColor(0xF7, 0x8C, 0x00)
PURPLE     = RGBColor(0x6E, 0x40, 0xC9)
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GREY = RGBColor(0xC9, 0xD1, 0xD9)
BLUE_TEXT  = RGBColor(0x58, 0xA6, 0xFF)
GREEN_TEXT = RGBColor(0x3F, 0xB9, 0x50)
YELLOW     = RGBColor(0xF7, 0xC9, 0x48)


def rgb(r, g, b):
    return RGBColor(r, g, b)


def add_rect(slide, left, top, width, height, fill_color, alpha=None):
    shape = slide.shapes.add_shape(
        pptx.enum.shapes.MSO_SHAPE_TYPE.AUTO_SHAPE if False else 1,
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.fill.background()
    return shape


def add_text(slide, text, left, top, width, height,
             font_size=10, bold=False, color=WHITE,
             align=PP_ALIGN.LEFT, wrap=True, italic=False):
    txBox = slide.shapes.add_textbox(
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    txBox.word_wrap = wrap
    tf = txBox.text_frame
    tf.word_wrap = wrap
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return txBox


def add_multiline(slide, lines, left, top, width, height,
                  font_size=9, color=LIGHT_GREY, spacing=1.0):
    txBox = slide.shapes.add_textbox(
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    txBox.word_wrap = True
    tf = txBox.text_frame
    tf.word_wrap = True
    first = True
    for line_data in lines:
        if isinstance(line_data, str):
            text, fc, bold, italic = line_data, color, False, False
        else:
            text   = line_data.get("text", "")
            fc     = line_data.get("color", color)
            bold   = line_data.get("bold", False)
            italic = line_data.get("italic", False)

        if first:
            p = tf.paragraphs[0]
            first = False
        else:
            p = tf.add_paragraph()

        run = p.add_run()
        run.text = text
        run.font.size = Pt(font_size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = fc
    return txBox


# ══════════════════════════════════════════════════════════════════════════════
# BACKGROUND
# ══════════════════════════════════════════════════════════════════════════════
add_rect(slide, 0, 0, 13.33, 7.5, DARK_BG)

# ══════════════════════════════════════════════════════════════════════════════
# HEADER BAR
# ══════════════════════════════════════════════════════════════════════════════
add_rect(slide, 0, 0, 13.33, 0.65, HEADER_BG)
add_text(slide, "UTILITICS — Azure Databricks Data Sharing Platform  |  Part 1 Complete + Part 2 Azure In Progress",
         0.15, 0.08, 10, 0.5, font_size=14, bold=True, color=WHITE)
add_text(slide, "Part 1 DONE  |  Part 2 IN PROGRESS",
         10.5, 0.08, 2.7, 0.5, font_size=11, bold=True, color=YELLOW, align=PP_ALIGN.RIGHT)

# ══════════════════════════════════════════════════════════════════════════════
# ROW 1  — three panels: Data Sources | Medallion Pipeline | Delta Sharing
# ══════════════════════════════════════════════════════════════════════════════

# ── Panel 1: Data Sources ─────────────────────────────────────────────────────
add_rect(slide, 0.1, 0.75, 3.1, 2.15, PANEL_BG)
add_rect(slide, 0.1, 0.75, 3.1, 0.3, GREEN)
add_text(slide, "① DATA SOURCES  (3 formats)", 0.15, 0.77, 3.0, 0.28,
         font_size=9, bold=True, color=WHITE)

add_multiline(slide, [
    {"text": "Time Series (CSV file drops)", "color": BLUE_TEXT, "bold": True},
    {"text": "  500 meters x 7 days x 48 readings"},
    {"text": "  168,000 records  |  data/lake/sources/timeseries_csv/"},
    {"text": ""},
    {"text": "Network Asset Snapshot (SQLite)", "color": BLUE_TEXT, "bold": True},
    {"text": "  2,000 assets x 7 snapshots"},
    {"text": "  14,000 records  |  data/lake/sources/assets.db"},
    {"text": "  JDBC read in bronze  (org.xerial:sqlite-jdbc)"},
    {"text": ""},
    {"text": "Demand Forecasts (Parquet)", "color": BLUE_TEXT, "bold": True},
    {"text": "  500 forecasts x 168 hourly points"},
    {"text": "  84,000 records  |  data/lake/sources/files_parquet/"},
], 0.15, 1.08, 3.0, 1.75, font_size=8)

# ── Panel 2: Medallion Pipeline ───────────────────────────────────────────────
add_rect(slide, 3.3, 0.75, 5.1, 2.15, PANEL_BG)
add_rect(slide, 3.3, 0.75, 5.1, 0.3, GREEN)
add_text(slide, "② MEDALLION PIPELINE  (PySpark 3.5.1 + Delta Lake 3.1.0)",
         3.35, 0.77, 5.0, 0.28, font_size=9, bold=True, color=WHITE)

add_multiline(slide, [
    {"text": "01  Generate Data", "color": YELLOW, "bold": True},
    {"text": "    Synthetic meter + asset + forecast  ->  266K records (CSV/SQLite/Parquet)"},
    {"text": "02  Bronze Ingestion", "color": YELLOW, "bold": True},
    {"text": "    Raw -> Bronze Delta  |  168K + 14K + 84K  |  Match: 100%"},
    {"text": "03  Silver Transformation", "color": YELLOW, "bold": True},
    {"text": "    Dedup + validation + derived cols  |  Drop rate: 0.0%  PASS"},
    {"text": "04  Gold Aggregation", "color": YELLOW, "bold": True},
    {"text": "    4 Gold tables: daily_meter / regional_demand /"},
    {"text": "    network_assets / forecast_summary  ->  6,760 records"},
    {"text": "05  Validation", "color": YELLOW, "bold": True},
    {"text": "    19/19 checks PASS  |  Full run: ~234 seconds"},
], 3.35, 1.08, 4.95, 1.75, font_size=8)

# ── Panel 3: Delta Sharing ────────────────────────────────────────────────────
add_rect(slide, 8.5, 0.75, 4.73, 2.15, PANEL_BG)
add_rect(slide, 8.5, 0.75, 4.73, 0.3, GREEN)
add_text(slide, "③ DELTA SHARING",
         8.55, 0.77, 4.6, 0.28, font_size=9, bold=True, color=WHITE)

add_multiline(slide, [
    {"text": "Option A — Python Simulation  [DONE]", "color": GREEN_TEXT, "bold": True},
    {"text": "  Provider shares 4 Gold tables with Vendor"},
    {"text": "  Vendor receives + analyses regional demand & assets"},
    {"text": "  Bi-directional: 5 forecast responses written back"},
    {"text": "  NORTH region highest peak: 670 kWh"},
    {"text": "  52.8% of assets flagged for maintenance"},
    {"text": ""},
    {"text": "Option B — Unity Catalog (Part 2)", "color": BLUE_TEXT, "bold": True},
    {"text": "  Azure Databricks Premium + Unity Catalog provisioned"},
    {"text": "  Real Delta Sharing via Unity Catalog — NEXT STEP"},
    {"text": "  Vendor gets read-only token  |  Full audit trail"},
], 8.55, 1.08, 4.6, 1.75, font_size=8)

# ══════════════════════════════════════════════════════════════════════════════
# ROW 2 — CI/CD | ChatOps | Part 2 Azure
# ══════════════════════════════════════════════════════════════════════════════

# ── Panel 4: CI/CD ────────────────────────────────────────────────────────────
add_rect(slide, 0.1, 3.05, 4.05, 2.2, PANEL_BG)
add_rect(slide, 0.1, 3.05, 4.05, 0.3, ORANGE)
add_text(slide, "④ CI/CD + NIGHTLY BUILD  [DONE]",
         0.15, 3.07, 3.95, 0.28, font_size=9, bold=True, color=WHITE)

add_multiline(slide, [
    {"text": "GitHub Actions CI  (.github/workflows/ci.yml)", "color": YELLOW, "bold": True},
    {"text": "  check_syntax.py: 25 files, syntax + 6 secret patterns"},
    {"text": "  Skips: ci/, notebooks/databricks/, vendor delta-sharing"},
    {"text": ""},
    {"text": "GitHub Actions CD  (.github/workflows/cd.yml)", "color": YELLOW, "bold": True},
    {"text": "  Triggers after CI passes on push to main"},
    {"text": "  upload_notebooks.py -> 9 notebooks auto-deployed"},
    {"text": "  Requires: DATABRICKS_HOST + DATABRICKS_TOKEN secrets"},
    {"text": ""},
    {"text": "Nightly Build  (ci/nightly_build.py)", "color": YELLOW, "bold": True},
    {"text": "  Windows Task Scheduler at 2AM daily"},
    {"text": "  Runs CI + full pipeline -> WhatsApp alert via Twilio"},
    {"text": "  setup_nightly.bat registers scheduled task"},
], 0.15, 3.38, 3.95, 1.8, font_size=8)

# ── Panel 5: ChatOps ──────────────────────────────────────────────────────────
add_rect(slide, 4.25, 3.05, 5.1, 2.2, PANEL_BG)
add_rect(slide, 4.25, 3.05, 5.1, 0.3, ORANGE)
add_text(slide, "⑤ CHATOPS — WhatsApp Pipeline Control  [LIVE]",
         4.3, 3.07, 5.0, 0.28, font_size=9, bold=True, color=WHITE)

add_multiline(slide, [
    {"text": "Flow:  WhatsApp -> Twilio -> ngrok -> FastAPI -> Claude AI -> Pipeline",
     "color": BLUE_TEXT, "bold": True},
    {"text": ""},
    {"text": "automation/webhook_server.py  (FastAPI)", "color": YELLOW, "bold": True},
    {"text": "  POST /webhook  receives Twilio WhatsApp messages"},
    {"text": "  Agentic loop: Claude calls tools, executes, replies"},
    {"text": ""},
    {"text": "automation/claude_tools.py  (7 tools)", "color": YELLOW, "bold": True},
    {"text": "  run_pipeline    run_step      run_sharing"},
    {"text": "  get_status      get_runtime   set_runtime"},
    {"text": "  run_ci  (NEW: triggers CI check from WhatsApp)"},
    {"text": ""},
    {"text": "Tested live: status, run pipeline, run_ci  [DONE]",
     "color": GREEN_TEXT, "bold": True},
], 4.3, 3.38, 4.95, 1.8, font_size=8)

# ── Panel 6: Part 2 Azure ─────────────────────────────────────────────────────
add_rect(slide, 9.45, 3.05, 3.78, 2.2, PANEL_BG)
add_rect(slide, 9.45, 3.05, 3.78, 0.3, PURPLE)
add_text(slide, "⑥ PART 2 — AZURE CLOUD  [IN PROGRESS]",
         9.5, 3.07, 3.68, 0.28, font_size=9, bold=True, color=WHITE)

add_multiline(slide, [
    {"text": "+ Resource group + ADLS Gen2 provisioned", "color": GREEN_TEXT, "bold": True},
    {"text": "+ Databricks Premium + Unity Catalog enabled", "color": GREEN_TEXT, "bold": True},
    {"text": "+ Cluster running  (Standard_DS3_v2)", "color": GREEN_TEXT, "bold": True},
    {"text": "+ 9 notebooks deployed via CI/CD  (abfss://)", "color": GREEN_TEXT, "bold": True},
    {"text": "+ Secret scope configured", "color": GREEN_TEXT, "bold": True},
    {"text": "o Run full pipeline 01->05 on Databricks", "color": ORANGE, "bold": True},
    {"text": "o Register Delta tables in Unity Catalog", "color": ORANGE},
    {"text": "o Real Delta Sharing — Phase 8", "color": ORANGE},
    {"text": "o Databricks Workflows — Phase 9", "color": ORANGE},
    {"text": "o Azure Functions ChatOps — Phase 10", "color": ORANGE},
    {"text": "Budget: ~$5-8/month  |  Production: ~$80-120/month",
     "color": LIGHT_GREY, "italic": True},
], 9.5, 3.38, 3.68, 1.8, font_size=8)

# ══════════════════════════════════════════════════════════════════════════════
# ROW 3 — Stats bar + Tech Stack + Key Fixes
# ══════════════════════════════════════════════════════════════════════════════

# Stats strip
add_rect(slide, 0.1, 5.38, 13.13, 0.52, rgb(0x0D, 0x1F, 0x0D))

stats = [
    ("266K",   "Raw Records"),
    ("19/19",  "Checks Passed"),
    ("6,760",  "Gold Records"),
    ("~234s",  "Full Pipeline"),
    ("7",      "Claude Tools"),
    ("9",      "Notebooks\nDeployed"),
    ("CI/CD",  "GitHub Actions"),
    ("LIVE",   "ChatOps"),
]
col_w = 13.13 / len(stats)
for i, (val, label) in enumerate(stats):
    x = 0.1 + i * col_w
    add_text(slide, val,   x, 5.4,  col_w, 0.25, font_size=12, bold=True,
             color=GREEN_TEXT, align=PP_ALIGN.CENTER)
    add_text(slide, label, x, 5.62, col_w, 0.2,  font_size=7,
             color=LIGHT_GREY, align=PP_ALIGN.CENTER)

# ── Tech stack + key fixes row ────────────────────────────────────────────────
add_rect(slide, 0.1, 6.0, 8.5, 1.38, PANEL_BG)
add_text(slide, "TECH STACK", 0.15, 6.02, 8.4, 0.22,
         font_size=8, bold=True, color=BLUE_TEXT)

tech = [
    "Python 3.11  |  PySpark 3.5.1  |  Delta Lake 3.1.0  |  SQLite + SQLAlchemy  |  sqlite-jdbc 3.44.1.0",
    "FastAPI + Uvicorn  |  Twilio WhatsApp API  |  Anthropic Claude API (tool use)  |  ngrok  |  python-dotenv",
    "Azure Databricks Premium  |  Unity Catalog  |  ADLS Gen2 (abfss://)  |  GitHub Actions CI/CD",
    "Windows 11 (24 GB RAM, 1 TB SSD)  |  Java 11  |  Hadoop winutils  |  venv",
]
for i, line in enumerate(tech):
    add_text(slide, line, 0.15, 6.25 + i * 0.27, 8.4, 0.26,
             font_size=7.5, color=LIGHT_GREY)

# Key fixes panel
add_rect(slide, 8.7, 6.0, 4.53, 1.38, PANEL_BG)
add_text(slide, "KEY WINDOWS FIXES", 8.75, 6.02, 4.4, 0.22,
         font_size=8, bold=True, color=ORANGE)

fixes = [
    "PYSPARK_PYTHON = sys.executable  (App Execution Alias bypass)",
    "data/lake/ path  (ACL-locked lake/ directories — Python creates with write ACL)",
    "CSV via toPandas().to_csv()  (replaces Parquet for time series drop)",
    "mode(overwrite) + shutil.move()  (Parquet overwrite on Windows)",
    "sys.stdout.reconfigure(encoding='utf-8')  (Unicode on Windows)",
    "SQLite JDBC type casts  (DateType/DoubleType mismatch)",
]
for i, line in enumerate(fixes):
    add_text(slide, f"- {line}", 8.75, 6.22 + i * 0.185, 4.4, 0.2,
             font_size=7, color=LIGHT_GREY)

# ── Footer ────────────────────────────────────────────────────────────────────
add_rect(slide, 0, 7.35, 13.33, 0.15, HEADER_BG)
add_text(slide,
         "Utilitics Data Sharing POC  |  Part 1 Complete — Part 2: Azure Cloud In Progress  |  2026-04-08",
         0.15, 7.35, 13.0, 0.15, font_size=7, color=WHITE, align=PP_ALIGN.CENTER)

# ── Save ──────────────────────────────────────────────────────────────────────
out = "C:/Projects/databricks-poc/Utilitics_OnePager.pptx"
prs.save(out)
print(f"Saved: {out}")
