"""
Waterfall Methodology - Industry Standard SDLC
Multi-slide PowerPoint presentation
Output: databricks-poc/methodology/Waterfall_Methodology.pptx
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import os

# ── Presentation Setup ────────────────────────────────────────────────────────
prs = Presentation()
prs.slide_width  = Inches(13.33)
prs.slide_height = Inches(7.5)

# ── Color Palette ─────────────────────────────────────────────────────────────
NAVY       = RGBColor(0x00, 0x2B, 0x5C)
BLUE       = RGBColor(0x00, 0x5B, 0xB5)
LIGHT_BLUE = RGBColor(0x33, 0x99, 0xFF)
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GREY = RGBColor(0xF0, 0xF4, 0xF8)
MID_GREY   = RGBColor(0x8A, 0x9B, 0xAD)
DARK_GREY  = RGBColor(0x33, 0x33, 0x33)
ORANGE     = RGBColor(0xFF, 0x66, 0x00)
GREEN      = RGBColor(0x00, 0x99, 0x33)
TEAL       = RGBColor(0x00, 0x80, 0x80)
RED        = RGBColor(0xCC, 0x00, 0x00)

PHASE_COLORS = [
    RGBColor(0x00, 0x5B, 0xB5),  # Blue        - Phase 1 Requirements
    RGBColor(0x00, 0x80, 0x80),  # Teal        - Phase 2 Design
    RGBColor(0x33, 0x66, 0x99),  # Steel Blue  - Phase 3 Development
    RGBColor(0xFF, 0x66, 0x00),  # Orange      - Phase 4 Testing
    RGBColor(0x00, 0x99, 0x33),  # Green       - Phase 5 Deployment
    RGBColor(0x60, 0x30, 0x99),  # Purple      - Phase 6 ELS/BAU
]

# ── Helper Functions ──────────────────────────────────────────────────────────
def blank_slide():
    return prs.slides.add_slide(prs.slide_layouts[6])

def add_rect(slide, left, top, width, height, fill_color,
             line_color=None, line_width=0):
    shape = slide.shapes.add_shape(
        1,
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if line_color:
        shape.line.color.rgb = line_color
        shape.line.width = Pt(line_width)
    else:
        shape.line.fill.background()
    return shape

def add_text(slide, text, left, top, width, height,
             font_size=10, bold=False, color=WHITE,
             align=PP_ALIGN.LEFT, italic=False, wrap=True):
    txBox = slide.shapes.add_textbox(
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    txBox.word_wrap = wrap
    tf = txBox.text_frame
    tf.word_wrap = wrap
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return txBox

def add_multiline(slide, lines, left, top, width, height,
                  font_size=9, default_color=DARK_GREY):
    txBox = slide.shapes.add_textbox(
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    txBox.word_wrap = True
    tf = txBox.text_frame
    tf.word_wrap = True
    first = True
    for line_data in lines:
        if isinstance(line_data, str):
            text, fc, bold, italic = line_data, default_color, False, False
        else:
            text   = line_data.get("text", "")
            fc     = line_data.get("color", default_color)
            bold   = line_data.get("bold", False)
            italic = line_data.get("italic", False)
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        run = p.add_run()
        run.text = text
        run.font.size = Pt(font_size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = fc
    return txBox

def add_header(slide, title, subtitle=None):
    add_rect(slide, 0, 0, 13.33, 7.5, LIGHT_GREY)
    add_rect(slide, 0, 0, 13.33, 1.0, NAVY)
    add_text(slide, title, 0.3, 0.08, 10, 0.52,
             font_size=20, bold=True, color=WHITE)
    if subtitle:
        add_text(slide, subtitle, 0.3, 0.6, 12.5, 0.35,
                 font_size=10, color=LIGHT_BLUE, italic=True)
    add_rect(slide, 0, 7.3, 13.33, 0.2, NAVY)
    add_text(slide,
             "Waterfall Methodology  |  Industry Standard SDLC  |  Confidential",
             0.2, 7.32, 10, 0.18, font_size=7, color=WHITE)
    add_text(slide, "2026", 11.5, 7.32, 1.8, 0.18,
             font_size=7, color=WHITE, align=PP_ALIGN.RIGHT)


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 1 — TITLE
# ══════════════════════════════════════════════════════════════════════════════
s1 = blank_slide()
add_rect(s1, 0, 0, 13.33, 7.5, NAVY)
add_rect(s1, 0, 3.45, 13.33, 0.06, ORANGE)

# Cascading waterfall blocks (decorative)
cascade_labels  = ["Plan", "Design", "Build", "Test", "Deploy", "Maintain"]
for i, (label, color) in enumerate(zip(cascade_labels, PHASE_COLORS)):
    x = 0.9 + i * 1.7
    y = 4.55 + i * 0.28
    add_rect(s1, x, y, 1.45, 0.3, color)
    add_text(s1, label, x + 0.05, y + 0.05, 1.35, 0.22,
             font_size=9, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

add_text(s1, "WATERFALL METHODOLOGY",
         0.8, 0.9, 11.7, 1.1,
         font_size=42, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
add_text(s1, "Industry Standard Software Development Life Cycle",
         0.8, 2.1, 11.7, 0.55,
         font_size=18, color=LIGHT_BLUE, align=PP_ALIGN.CENTER)
add_text(s1, "A structured, sequential, phase-gated approach to software delivery",
         0.8, 2.75, 11.7, 0.4,
         font_size=13, italic=True, color=MID_GREY, align=PP_ALIGN.CENTER)

add_rect(s1, 0, 7.2, 13.33, 0.3, RGBColor(0x00, 0x1A, 0x40))
add_text(s1, "Confidential  |  2026  |  Industry Standard Framework",
         0.5, 7.22, 12.5, 0.25,
         font_size=8, color=MID_GREY, align=PP_ALIGN.CENTER)


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 2 — OVERVIEW
# ══════════════════════════════════════════════════════════════════════════════
s2 = blank_slide()
add_header(s2, "Waterfall SDLC — Phase Overview",
           "Sequential, phase-gated delivery model — each phase signed off before the next begins")

phases = [
    ("1", "Planning &\nRequirements", "• BID / Feasibility\n• BRD / Use Cases\n• Requirements Freeze\n• RTM"),
    ("2", "System\nDesign",           "• HLD / LLD\n• Architecture\n• ARB Submission\n• Test Strategy"),
    ("3", "Development\n& Build",     "• Coding\n• Code Review\n• Unit Testing\n• Build Artifacts"),
    ("4", "Testing\nSIT & UAT",       "• SIT / Performance\n• Security Testing\n• UAT\n• UAT Sign-off"),
    ("5", "Deployment\n& Go-Live",    "• CAB Approval\n• Cutover Runbook\n• Go-Live\n• Smoke Test"),
    ("6", "ELS &\nMaintenance",       "• Hypercare (ELS)\n• PIR\n• BAU Handover\n• Project Closure"),
]

box_w = 1.95
box_h = 3.8
start_x = 0.2

for i, (num, title, desc) in enumerate(phases):
    x     = start_x + i * (box_w + 0.13)
    color = PHASE_COLORS[i]
    add_rect(s2, x, 1.15, box_w, box_h, WHITE, line_color=color, line_width=1.5)
    add_rect(s2, x, 1.15, box_w, 0.7, color)
    add_text(s2, f"Phase {num}", x + 0.1, 1.17, box_w - 0.2, 0.25,
             font_size=9, bold=True, color=WHITE)
    add_text(s2, title, x + 0.1, 1.4, box_w - 0.2, 0.42,
             font_size=11, bold=True, color=WHITE)
    if i < len(phases) - 1:
        ax = x + box_w + 0.02
        add_text(s2, "→", ax, 2.6, 0.14, 0.3,
                 font_size=16, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    add_multiline(s2, desc.split("\n"),
                  x + 0.1, 1.93, box_w - 0.2, box_h - 0.9,
                  font_size=9, default_color=DARK_GREY)

# Key principle note
add_rect(s2, 0.2, 5.1, 12.9, 0.5, RGBColor(0xE8, 0xF0, 0xFE))
add_text(s2,
         "Key Principle: Each phase must be fully completed and formally signed off before the next phase begins. "
         "Any changes after Requirements Freeze must go through formal Change Control.",
         0.35, 5.14, 12.6, 0.42, font_size=9, color=NAVY, italic=True)

# Gates bar
add_rect(s2, 0.2, 5.72, 12.9, 0.52, RGBColor(0xFF, 0xF3, 0xE0))
add_text(s2, "Gates: ", 0.35, 5.76, 1.0, 0.4,
         font_size=9, bold=True, color=ORANGE)
add_text(s2,
         "Gate 1: Requirements Sign-off  →  Gate 2: ARB + Design Approval  →  "
         "Gate 3: Build Complete  →  Gate 4: SIT Pass + UAT Sign-off  →  "
         "Gate 5: CAB Approval (Go-Live)  →  Gate 6: ELS Sign-off + BAU Handover",
         1.25, 5.76, 11.7, 0.4, font_size=8.5, color=DARK_GREY)


# ══════════════════════════════════════════════════════════════════════════════
# SLIDES 3–8 — ONE PER PHASE
# ══════════════════════════════════════════════════════════════════════════════

phase_data = [
    # (slide_title, subtitle, color_index, left_title, left_lines, mid_title, mid_lines, entry_lines, exit_title, exit_lines)
    {
        "title":    "Phase 1 — Planning & Requirements",
        "subtitle": "Establish project foundation and capture all business needs before design begins",
        "ci":       0,
        "lt":       "Key Activities",
        "ll": [
            {"text": "Project Initiation", "color": PHASE_COLORS[0], "bold": True},
            "  • Business case / BID preparation",
            "  • Feasibility study (technical, financial, operational)",
            "  • Project charter & stakeholder identification",
            "  • Resource planning & cost estimation",
            {"text": "Requirements Gathering", "color": PHASE_COLORS[0], "bold": True},
            "  • Stakeholder interviews & workshops",
            "  • Business Requirements Document (BRD)",
            "  • Functional Specification Document (FSD)",
            "  • Use case definition & documentation",
            "  • Non-Functional Requirements (NFR)",
            {"text": "Requirements Freeze", "color": PHASE_COLORS[0], "bold": True},
            "  • Requirements review & walkthrough",
            "  • Formal sign-off by business stakeholders",
            "  • Change control activated post-freeze",
            {"text": "Traceability", "color": PHASE_COLORS[0], "bold": True},
            "  • Requirements Traceability Matrix (RTM) created",
        ],
        "mt": "Key Documents & Deliverables",
        "ml": [
            {"text": "BID Document", "color": PHASE_COLORS[0], "bold": True},
            "  Business case, cost-benefit, ROI analysis",
            {"text": "Project Charter", "color": PHASE_COLORS[0], "bold": True},
            "  Scope, objectives, sponsor approval",
            {"text": "BRD (Business Requirements Doc)", "color": PHASE_COLORS[0], "bold": True},
            "  Full functional & non-functional requirements",
            {"text": "FSD (Functional Specification Doc)", "color": PHASE_COLORS[0], "bold": True},
            "  Detailed functional scenario specifications",
            {"text": "Use Case Document", "color": PHASE_COLORS[0], "bold": True},
            "  Actors, flows, pre/post conditions",
            {"text": "RTM (Requirements Traceability Matrix)", "color": PHASE_COLORS[0], "bold": True},
            "  Maps requirements to test cases",
            {"text": "Risk Register", "color": PHASE_COLORS[0], "bold": True},
            "  Identified risks, likelihood, mitigation",
            {"text": "Project Plan / RACI", "color": PHASE_COLORS[0], "bold": True},
            "  Milestones, resources, responsibilities",
        ],
        "el": [
            "  ✔  Business need / opportunity identified",
            "  ✔  Executive sponsor assigned",
            "  ✔  Initial budget approved",
            "  ✔  Key stakeholders identified",
        ],
        "xt": "Exit Criteria — Gate 1",
        "xl": [
            "  ✔  BRD & FSD reviewed and approved",
            "  ✔  Requirements signed off by business",
            "  ✔  RTM baselined",
            "  ✔  Project plan approved",
            "  ✔  Risks documented and accepted",
            "  ✔  Change control process activated",
        ],
    },
    {
        "title":    "Phase 2 — System Design",
        "subtitle": "Translate requirements into technical architecture and detailed design blueprints",
        "ci":       1,
        "lt":       "Key Activities",
        "ll": [
            {"text": "High Level Design (HLD)", "color": PHASE_COLORS[1], "bold": True},
            "  • System architecture & component design",
            "  • Technology stack selection",
            "  • Integration design & data flows",
            "  • Infrastructure & hosting design",
            {"text": "Low Level Design (LLD)", "color": PHASE_COLORS[1], "bold": True},
            "  • Database schema design (ERD)",
            "  • Module / class level design",
            "  • API specifications",
            "  • Security design",
            {"text": "ARB Submission", "color": PHASE_COLORS[1], "bold": True},
            "  • Architecture Review Board presentation",
            "  • Technical standards compliance check",
            "  • Security & risk review",
            "  • ARB approval required before build starts",
            {"text": "Test Strategy", "color": PHASE_COLORS[1], "bold": True},
            "  • Overall test approach defined",
            "  • Test environments planned & provisioned",
        ],
        "mt": "Key Documents & Deliverables",
        "ml": [
            {"text": "HLD Document", "color": PHASE_COLORS[1], "bold": True},
            "  Architecture diagrams, component overview",
            {"text": "LLD Document", "color": PHASE_COLORS[1], "bold": True},
            "  Detailed technical specifications",
            {"text": "ARB Submission Pack", "color": PHASE_COLORS[1], "bold": True},
            "  Architecture proposal for governance board",
            {"text": "Database Design Document", "color": PHASE_COLORS[1], "bold": True},
            "  ERD, schema, data dictionary",
            {"text": "Interface Design Document (IDD)", "color": PHASE_COLORS[1], "bold": True},
            "  API specs, integration contracts",
            {"text": "Security Design Document", "color": PHASE_COLORS[1], "bold": True},
            "  Authentication, authorisation, encryption",
            {"text": "Test Strategy Document", "color": PHASE_COLORS[1], "bold": True},
            "  Test scope, approach, environments",
            {"text": "Updated RTM", "color": PHASE_COLORS[1], "bold": True},
            "  Design elements mapped to requirements",
        ],
        "el": [
            "  ✔  Gate 1 passed — BRD signed off",
            "  ✔  RTM baselined",
            "  ✔  Architects assigned",
            "  ✔  Dev environment available",
        ],
        "xt": "Exit Criteria — Gate 2",
        "xl": [
            "  ✔  HLD & LLD reviewed and approved",
            "  ✔  ARB approval obtained",
            "  ✔  Security design signed off",
            "  ✔  Test strategy approved",
            "  ✔  Test environments provisioned",
            "  ✔  RTM updated with design traceability",
        ],
    },
    {
        "title":    "Phase 3 — Development & Build",
        "subtitle": "Code, review and unit test all components per the approved Low Level Design",
        "ci":       2,
        "lt":       "Key Activities",
        "ll": [
            {"text": "Development Setup", "color": PHASE_COLORS[2], "bold": True},
            "  • Version control branching strategy (Git)",
            "  • Coding standards & conventions defined",
            "  • CI/CD pipeline setup",
            {"text": "Build & Code", "color": PHASE_COLORS[2], "bold": True},
            "  • Feature development per LLD",
            "  • Database objects (tables, views, procedures)",
            "  • API & integration development",
            "  • Configuration & environment setup",
            {"text": "Code Quality", "color": PHASE_COLORS[2], "bold": True},
            "  • Peer code review",
            "  • Static code analysis",
            "  • Coding standards compliance check",
            {"text": "Unit Testing", "color": PHASE_COLORS[2], "bold": True},
            "  • Developer-written unit tests",
            "  • Test coverage targets met",
            "  • Bug fixes and re-testing",
            {"text": "Build Management", "color": PHASE_COLORS[2], "bold": True},
            "  • Release build packaging",
            "  • Deployment scripts prepared",
        ],
        "mt": "Key Documents & Deliverables",
        "ml": [
            {"text": "Source Code", "color": PHASE_COLORS[2], "bold": True},
            "  Version-controlled in Git repository",
            {"text": "Unit Test Results", "color": PHASE_COLORS[2], "bold": True},
            "  Coverage reports, pass/fail logs",
            {"text": "Code Review Records", "color": PHASE_COLORS[2], "bold": True},
            "  Pull requests, review comments",
            {"text": "Build Artifacts", "color": PHASE_COLORS[2], "bold": True},
            "  Compiled binaries, packages, scripts",
            {"text": "Deployment Scripts", "color": PHASE_COLORS[2], "bold": True},
            "  DB migration scripts, config scripts",
            {"text": "Technical Documentation", "color": PHASE_COLORS[2], "bold": True},
            "  Code documentation, README files",
            {"text": "Updated RTM", "color": PHASE_COLORS[2], "bold": True},
            "  Components mapped to requirements",
            {"text": "Build Release Notes", "color": PHASE_COLORS[2], "bold": True},
            "  What has been built, known issues",
        ],
        "el": [
            "  ✔  Gate 2 passed — Design approved",
            "  ✔  ARB approval in place",
            "  ✔  Dev environment provisioned",
            "  ✔  Development team assigned",
        ],
        "xt": "Exit Criteria — Gate 3",
        "xl": [
            "  ✔  All features built per LLD",
            "  ✔  Code reviews completed",
            "  ✔  Unit tests passed (coverage met)",
            "  ✔  No critical/high open defects",
            "  ✔  Build artifacts deployable",
            "  ✔  SIT environment ready",
        ],
    },
    {
        "title":    "Phase 4 — Testing (SIT & UAT)",
        "subtitle": "Validate the system against requirements through structured, evidence-based testing",
        "ci":       3,
        "lt":       "SIT — System Integration Testing",
        "ll": [
            {"text": "Purpose", "color": PHASE_COLORS[3], "bold": True},
            "  Verify all integrated components work",
            "  together as a complete system",
            {"text": "Activities", "color": PHASE_COLORS[3], "bold": True},
            "  • Execute SIT test cases (functional)",
            "  • Integration & end-to-end testing",
            "  • Performance & load testing",
            "  • Security & penetration testing",
            "  • Regression testing",
            "  • Defect logging, triage & fix cycle",
            "  • Re-testing of resolved defects",
            {"text": "Managed By", "color": PHASE_COLORS[3], "bold": True},
            "  QA / Test team",
            {"text": "SIT Sign-off", "color": PHASE_COLORS[3], "bold": True},
            "  • All P1/P2 defects resolved",
            "  • Test evidence documented",
            "  • SIT sign-off by Test Lead",
        ],
        "mt": "UAT — User Acceptance Testing",
        "ml": [
            {"text": "Purpose", "color": PHASE_COLORS[3], "bold": True},
            "  Business validates the system meets",
            "  requirements in a real-world scenario",
            {"text": "Activities", "color": PHASE_COLORS[3], "bold": True},
            "  • Business executes UAT test scripts",
            "  • Test against BRD & use cases",
            "  • Business process walkthroughs",
            "  • Data validation & output verification",
            "  • Defect logging & priority assignment",
            "  • Fix, re-deploy, re-test cycle",
            "  • PAT (Parallel Acceptance Testing)",
            {"text": "Managed By", "color": PHASE_COLORS[3], "bold": True},
            "  Business stakeholders / Product Owner",
            {"text": "UAT Sign-off", "color": PHASE_COLORS[3], "bold": True},
            "  • Written sign-off from business",
            "  • All P1/P2 UAT defects resolved",
            "  • Formal UAT completion certificate",
        ],
        "el": [
            "  ✔  Gate 3 passed — Build complete",
            "  ✔  SIT environment provisioned",
            "  ✔  Test cases written & reviewed",
            "  ✔  Test data prepared",
        ],
        "xt": "Exit Criteria — Gate 4",
        "xl": [
            "  ✔  SIT completed & signed off",
            "  ✔  UAT completed & signed off",
            "  ✔  No open P1/P2 defects",
            "  ✔  Performance benchmarks met",
            "  ✔  Security testing passed",
            "  ✔  RTM — 100% requirements covered",
        ],
    },
    {
        "title":    "Phase 5 — Deployment & Go-Live",
        "subtitle": "Controlled release to production following formal governance and CAB approval",
        "ci":       4,
        "lt":       "Key Activities",
        "ll": [
            {"text": "Pre-Deployment", "color": PHASE_COLORS[4], "bold": True},
            "  • CAB submission & approval",
            "  • Production environment readiness check",
            "  • RMAN / database backup",
            "  • Rollback plan documented & approved",
            "  • Stakeholder communication sent",
            "  • Scheduler / Control-M jobs on hold",
            {"text": "Cutover Execution", "color": PHASE_COLORS[4], "bold": True},
            "  • Execute cutover runbook step-by-step",
            "  • Deploy DB objects & scripts",
            "  • Deploy application components",
            "  • Configure integrations & file transfers",
            "  • Verify each step before proceeding",
            {"text": "Go-Live Verification", "color": PHASE_COLORS[4], "bold": True},
            "  • Smoke test in production",
            "  • End-to-end sanity check",
            "  • Checkpoint email to stakeholders",
            "  • Release batch/scheduler jobs",
        ],
        "mt": "Key Documents & Deliverables",
        "ml": [
            {"text": "CAB Change Request", "color": PHASE_COLORS[4], "bold": True},
            "  Formal request for production change",
            {"text": "Cutover Runbook", "color": PHASE_COLORS[4], "bold": True},
            "  Step-by-step deployment instructions",
            {"text": "Rollback Plan", "color": PHASE_COLORS[4], "bold": True},
            "  Reversion steps if go-live fails",
            {"text": "Production Backup Evidence", "color": PHASE_COLORS[4], "bold": True},
            "  RMAN backup confirmation",
            {"text": "Deployment Log", "color": PHASE_COLORS[4], "bold": True},
            "  Record of all steps executed & outcomes",
            {"text": "Go-Live Communication", "color": PHASE_COLORS[4], "bold": True},
            "  Stakeholder notification emails",
            {"text": "Smoke Test Results", "color": PHASE_COLORS[4], "bold": True},
            "  Production verification evidence",
            {"text": "Release Notes", "color": PHASE_COLORS[4], "bold": True},
            "  What was deployed & known issues",
        ],
        "el": [
            "  ✔  Gate 4 passed — UAT signed off",
            "  ✔  CAB approval granted",
            "  ✔  Production backup taken",
            "  ✔  Rollback plan approved",
        ],
        "xt": "Exit Criteria — Gate 5",
        "xl": [
            "  ✔  All deployment steps completed",
            "  ✔  Smoke tests passed in production",
            "  ✔  Batch jobs released & running",
            "  ✔  No critical production incidents",
            "  ✔  Stakeholders notified of go-live",
            "  ✔  Deployment log completed & filed",
        ],
    },
    {
        "title":    "Phase 6 — ELS, BAU & Project Closure",
        "subtitle": "Stabilise the system, transition to operations, and formally close the project",
        "ci":       5,
        "lt":       "ELS — Early Life Support (Hypercare)",
        "ll": [
            {"text": "What is ELS?", "color": PHASE_COLORS[5], "bold": True},
            "  Heightened support period immediately",
            "  after go-live (typically 2–4 weeks)",
            {"text": "Activities", "color": PHASE_COLORS[5], "bold": True},
            "  • Extended-hours or 24/7 monitoring",
            "  • Batch job monitoring & verification",
            "  • Incident triage & rapid response",
            "  • Performance monitoring",
            "  • Daily ELS status reports to stakeholders",
            "  • Critical defect resolution",
            "  • End-user support & guidance",
            {"text": "ELS Exit Criteria", "color": PHASE_COLORS[5], "bold": True},
            "  • System stable for agreed ELS period",
            "  • No open P1/P2 incidents",
            "  • All batch runs completed successfully",
            "  • Formal ELS sign-off obtained",
        ],
        "mt": "BAU Handover & Project Closure",
        "ml": [
            {"text": "BAU Transition", "color": PHASE_COLORS[5], "bold": True},
            "  • Knowledge transfer to BAU / ops team",
            "  • Runbooks & SOPs handed over",
            "  • Support model defined (L1/L2/L3)",
            "  • On-call rota established",
            "  • Monitoring & alerting configured",
            {"text": "Post Implementation Review (PIR)", "color": PHASE_COLORS[5], "bold": True},
            "  • Lessons learned workshop",
            "  • Benefits realisation assessment",
            "  • PIR report produced & distributed",
            {"text": "Project Closure", "color": PHASE_COLORS[5], "bold": True},
            "  • Change request formally closed",
            "  • Budget reconciliation completed",
            "  • Resources released",
            "  • Final project report produced",
            "  • Sponsor sign-off on project closure",
        ],
        "el": [
            "  ✔  Gate 5 passed — Go-live complete",
            "  ✔  System live in production",
            "  ✔  Support team briefed",
            "  ✔  Monitoring tools active",
        ],
        "xt": "Exit Criteria — Gate 6",
        "xl": [
            "  ✔  ELS sign-off obtained",
            "  ✔  BAU team fully transitioned",
            "  ✔  PIR report completed",
            "  ✔  Change request closed",
            "  ✔  Lessons learned documented",
            "  ✔  Sponsor project closure sign-off",
        ],
    },
]

for pd in phase_data:
    s = blank_slide()
    color = PHASE_COLORS[pd["ci"]]
    add_header(s, pd["title"], pd["subtitle"])

    # Left panel
    add_rect(s, 0.25, 1.1, 4.1, 5.6, WHITE, line_color=color, line_width=1)
    add_rect(s, 0.25, 1.1, 4.1, 0.35, color)
    add_text(s, pd["lt"], 0.35, 1.13, 3.9, 0.28, font_size=10, bold=True, color=WHITE)
    add_multiline(s, pd["ll"], 0.35, 1.52, 3.9, 5.1, font_size=9, default_color=DARK_GREY)

    # Mid panel
    add_rect(s, 4.55, 1.1, 4.1, 5.6, WHITE, line_color=color, line_width=1)
    add_rect(s, 4.55, 1.1, 4.1, 0.35, color)
    add_text(s, pd["mt"], 4.65, 1.13, 3.9, 0.28, font_size=10, bold=True, color=WHITE)
    add_multiline(s, pd["ml"], 4.65, 1.52, 3.9, 5.1, font_size=9, default_color=DARK_GREY)

    # Entry criteria
    add_rect(s, 8.85, 1.1, 4.23, 2.65, WHITE, line_color=GREEN, line_width=1)
    add_rect(s, 8.85, 1.1, 4.23, 0.35, GREEN)
    add_text(s, "Entry Criteria", 8.95, 1.13, 4.1, 0.28, font_size=10, bold=True, color=WHITE)
    add_multiline(s, pd["el"], 8.95, 1.52, 4.1, 2.1, font_size=9, default_color=DARK_GREY)

    # Exit criteria
    add_rect(s, 8.85, 3.85, 4.23, 2.85, WHITE, line_color=RED, line_width=1)
    add_rect(s, 8.85, 3.85, 4.23, 0.35, RED)
    add_text(s, pd["xt"], 8.95, 3.88, 4.1, 0.28, font_size=10, bold=True, color=WHITE)
    add_multiline(s, pd["xl"], 8.95, 4.27, 4.1, 2.3, font_size=9, default_color=DARK_GREY)


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 9 — KEY DOCUMENTS & GOVERNANCE
# ══════════════════════════════════════════════════════════════════════════════
s9 = blank_slide()
add_header(s9, "Key Documents, Gates & Governance",
           "Summary of deliverables, approval gates and governance bodies across all phases")

add_rect(s9, 0.25, 1.1, 8.5, 5.65, WHITE, line_color=NAVY, line_width=1)
add_rect(s9, 0.25, 1.1, 8.5, 0.35, NAVY)
add_text(s9, "Key Documents by Phase", 0.35, 1.13, 8.3, 0.28,
         font_size=11, bold=True, color=WHITE)

docs_by_phase = [
    (PHASE_COLORS[0], "Phase 1 — Requirements",
     "BID  •  Project Charter  •  BRD  •  FSD  •  Use Cases  •  RTM  •  Risk Register  •  Project Plan"),
    (PHASE_COLORS[1], "Phase 2 — Design",
     "HLD  •  LLD  •  ARB Submission  •  DB Design  •  IDD  •  Security Design  •  Test Strategy"),
    (PHASE_COLORS[2], "Phase 3 — Development",
     "Source Code  •  Unit Test Results  •  Code Review Records  •  Build Artifacts  •  Deploy Scripts"),
    (PHASE_COLORS[3], "Phase 4 — Testing",
     "SIT Test Plan  •  SIT Results  •  UAT Test Scripts  •  Defect Log  •  UAT Sign-off Certificate"),
    (PHASE_COLORS[4], "Phase 5 — Deployment",
     "CAB Change Request  •  Cutover Runbook  •  Rollback Plan  •  Deployment Log  •  Release Notes"),
    (PHASE_COLORS[5], "Phase 6 — ELS/BAU",
     "ELS Status Reports  •  ELS Sign-off  •  BAU Runbooks  •  PIR Report  •  Project Closure Report"),
]

for i, (col, phase, docs) in enumerate(docs_by_phase):
    y = 1.52 + i * 0.72
    add_rect(s9, 0.25, y, 2.2, 0.65, col)
    add_text(s9, phase, 0.35, y + 0.1, 2.1, 0.5, font_size=8, bold=True, color=WHITE)
    add_rect(s9, 2.45, y, 6.3, 0.65,
             RGBColor(0xF8, 0xF9, 0xFA), line_color=RGBColor(0xDD, 0xDD, 0xDD), line_width=0.5)
    add_text(s9, docs, 2.55, y + 0.14, 6.1, 0.48, font_size=8, color=DARK_GREY)

# Governance panel
add_rect(s9, 8.85, 1.1, 4.23, 5.65, WHITE, line_color=NAVY, line_width=1)
add_rect(s9, 8.85, 1.1, 4.23, 0.35, NAVY)
add_text(s9, "Governance Bodies", 8.95, 1.13, 4.1, 0.28,
         font_size=11, bold=True, color=WHITE)

gov_items = [
    (ORANGE, "ARB — Architecture Review Board",
     "Reviews & approves technical architecture\nbefore development begins"),
    (RED,    "CAB — Change Advisory Board",
     "Reviews & approves all production\nchanges before deployment"),
    (BLUE,   "PMO — Project Management Office",
     "Oversees project governance, reporting\nand milestone tracking"),
    (TEAL,   "Steering Committee",
     "Executive oversight, budget approval\nand escalation resolution"),
    (GREEN,  "Change Control Board",
     "Manages post-freeze requirement changes\nand impact assessment"),
]

for i, (col, title, desc) in enumerate(gov_items):
    y = 1.52 + i * 1.02
    add_rect(s9, 8.85, y, 4.23, 0.95,
             RGBColor(0xF8, 0xF9, 0xFA), line_color=col, line_width=1.5)
    add_rect(s9, 8.85, y, 0.08, 0.95, col)
    add_text(s9, title, 9.0, y + 0.04, 4.05, 0.28, font_size=9, bold=True, color=col)
    add_text(s9, desc,  9.0, y + 0.34, 4.05, 0.58, font_size=8, color=DARK_GREY)


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 10 — ROLES & RESPONSIBILITIES
# ══════════════════════════════════════════════════════════════════════════════
s10 = blank_slide()
add_header(s10, "Roles & Responsibilities",
           "Key stakeholders and their involvement across the Waterfall lifecycle")

roles = [
    ("Executive Sponsor",     NAVY,            "Approves BID & budget, provides strategic direction, resolves escalations, signs off project closure"),
    ("Project Manager",       PHASE_COLORS[0], "Manages plan, budget, risks and issues. Drives phase gates, stakeholder communications and reporting"),
    ("Business Analyst",      PHASE_COLORS[1], "Leads requirements gathering, produces BRD, FSD, Use Cases, RTM. Supports UAT"),
    ("Solution Architect",    PHASE_COLORS[1], "Produces HLD/LLD, presents ARB submission, governs technical design decisions"),
    ("Development Team",      PHASE_COLORS[2], "Builds all components per LLD, conducts code reviews, unit testing, prepares deployment artefacts"),
    ("QA / Test Team",        PHASE_COLORS[3], "Writes and executes SIT test cases, manages defect log, signs off SIT completion"),
    ("Business Users",        PHASE_COLORS[3], "Executes UAT test scripts, raises defects, provides formal UAT sign-off"),
    ("Security Team",         PHASE_COLORS[4], "Reviews security design, conducts penetration testing, signs off security clearance"),
    ("Infrastructure / DBA",  PHASE_COLORS[4], "Provisions environments, manages database, supports backup/restore and production deployment"),
    ("Change Manager",        PHASE_COLORS[5], "Manages change control, coordinates CAB submissions, handles post-freeze change requests"),
    ("BAU / Support Team",    ORANGE,          "Participates in ELS, receives knowledge transfer, owns operational support post-handover"),
]

col_w = 4.1
for i, (role, color, desc) in enumerate(roles):
    col = i % 3
    row = i // 3
    x = 0.2 + col * (col_w + 0.2)
    y = 1.1 + row * 1.45
    add_rect(s10, x, y, col_w, 1.35, WHITE, line_color=color, line_width=1)
    add_rect(s10, x, y, col_w, 0.32, color)
    add_text(s10, role, x + 0.1, y + 0.04, col_w - 0.2, 0.26,
             font_size=9, bold=True, color=WHITE)
    add_text(s10, desc, x + 0.1, y + 0.38, col_w - 0.2, 0.9,
             font_size=8, color=DARK_GREY, wrap=True)


# ══════════════════════════════════════════════════════════════════════════════
# SAVE
# ══════════════════════════════════════════════════════════════════════════════
out = "C:/Projects/databricks-poc/methodology/Waterfall_Methodology.pptx"
prs.save(out)
print(f"Saved: {out}")
